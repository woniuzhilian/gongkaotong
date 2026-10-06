import json, sys, re
sys.stdout.reconfigure(encoding='utf-8')

# 1. 从questions_raw.txt提取题号，按顺序对应id（修正正则：允许年份和补之间有空格）
with open(r'D:\应用程序开发\刷题\questions_raw.txt', 'r', encoding='utf-8') as f:
    text = f.read()

# 修正后的正则：允许年份和"补"之间有空格，允许标记前后有空格
pattern = re.compile(r'【\s*(\d{4}\s*补?)\s*-\s*(\d+)\s*】')
matches = pattern.findall(text)

# 建立id到题号的映射（id从1开始，按出现顺序）
id_to_qnum = {}
for i, (year, qnum) in enumerate(matches, start=1):
    year_clean = year.replace(' ', '')
    id_to_qnum[i] = f"{year_clean}-{qnum}"

print(f"共提取到 {len(matches)} 个题目标记")
print(f"id=311对应: {id_to_qnum.get(311, '?')}")
print(f"id=64对应: {id_to_qnum.get(64, '?')}")
print(f"id=552对应: {id_to_qnum.get(552, '?')}")

# 验证：找PDF45页对应的id范围
# 2019-28应该是第几个标记
count_2019_28 = 0
for i, (year, qnum) in enumerate(matches, start=1):
    year_clean = year.replace(' ', '')
    if year_clean == '2019' and qnum == '28':
        print(f"2019-28 对应 id={i}")
        break

# 2. 加载JSON数据
with open(r'D:\应用程序开发\刷题\all_questions_final.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

# 3. 筛选题干明确引用图片的题目
img_keywords = ['如图', '图所示', '图中', '下图', '上图', '图示', '见图', '图为', '图是']
img_questions = []
for q in data:
    qtext = q['question']
    if any(kw in qtext for kw in img_keywords):
        m = re.search(r'PDF第(\d+)页', qtext)
        page = m.group(1) if m else ''
        qclean = re.sub(r'【本题配图，PDF第\d+页】', '', qtext).strip()
        qshort = qclean[:60] + ('...' if len(qclean) > 60 else '')
        img_questions.append({
            'id': q['id'],
            '题号': id_to_qnum.get(q['id'], ''),
            '年份': q['year'],
            '小科目': q['smallSubject'],
            'PDF页码': page,
            '题干摘要': qshort,
            '配图': ''
        })

print(f"\n需配图题目共 {len(img_questions)} 道")

# 4. 生成Excel
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "需配图题目清单"

headers = ['序号', 'id', '题号', '年份', '小科目', 'PDF页码', '题干摘要', '配图']
ws.append(headers)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font_white = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font_white
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

for idx, q in enumerate(img_questions, start=1):
    row = [idx, q['id'], q['题号'], q['年份'], q['小科目'], q['PDF页码'], q['题干摘要'], q['配图']]
    ws.append(row)
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=idx+1, column=col)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if col in [1, 2, 4, 6]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 6
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 8
ws.column_dimensions['E'].width = 22
ws.column_dimensions['F'].width = 9
ws.column_dimensions['G'].width = 60
ws.column_dimensions['H'].width = 30
ws.freeze_panes = 'A2'
ws.row_dimensions[1].height = 25

output_path = r'D:\应用程序开发\刷题\需配图题目清单.xlsx'
wb.save(output_path)
print(f"\nExcel已保存: {output_path}")
