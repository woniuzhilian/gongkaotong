import json, sys, re
sys.stdout.reconfigure(encoding='utf-8')

# 1. 从questions_raw.txt提取题号，按顺序对应id
with open(r'D:\应用程序开发\刷题\questions_raw.txt', 'r', encoding='utf-8') as f:
    text = f.read()

pattern = re.compile(r'【\s*(\d{4}\s*补?)\s*-\s*(\d+)\s*】')
matches = list(pattern.finditer(text))
print(f"PDF中总标记数: {len(matches)}")

# 找到重复的2016-66的位置
dup_positions = []
for i, m in enumerate(matches):
    year = m.group(1).replace(' ', '')
    qnum = m.group(2)
    if year == '2016' and qnum == '66':
        dup_positions.append(i+1)  # 1-based
print(f"2016-66出现位置: {dup_positions}")

# 建立id到题号的映射
# JSON中去掉了第二个2016-66（假设是第819个匹配）
# id=1..818 对应匹配1..818
# id=819..1438 对应匹配820..1439（跳过第819个重复）
skip_pos = dup_positions[1] if len(dup_positions) > 1 else None
print(f"跳过重复位置: {skip_pos}")

id_to_qnum = {}
for qid in range(1, 1439):  # id 1-1438
    if qid < skip_pos:
        match_idx = qid - 1  # 0-based
    else:
        match_idx = qid  # 跳过一个，所以qid对应第qid+1个匹配（0-based为qid）
    m = matches[match_idx]
    year = m.group(1).replace(' ', '')
    qnum = m.group(2)
    id_to_qnum[qid] = f"{year}-{qnum}"

# 验证
print(f"\n验证:")
print(f"id=153: {id_to_qnum[153]}")
print(f"id=312: {id_to_qnum[312]}")
print(f"id=818: {id_to_qnum[818]}")
print(f"id=819: {id_to_qnum[819]}")

# 2. 加载JSON数据
with open(r'D:\应用程序开发\刷题\all_questions_final.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

# 3. 筛选题干明确引用图片的题目
img_keywords = ['如图', '图所示', '图中', '下图', '上图', '图示', '见图', '图为', '图是']
img_questions = []
for q in data:
    qtext = q['question']
    if any(kw in qtext for kw in img_keywords):
        m = re.search(r'PDF第(\d+)页', qtext)
        page = m.group(1) if m else ''
        qclean = re.sub(r'【本题配图，PDF第\d+页】', '', qtext).strip()
        qshort = qclean[:60] + ('...' if len(qclean) > 60 else '')
        img_questions.append({
            'id': q['id'],
            '题号': id_to_qnum.get(q['id'], ''),
            '年份': q['year'],
            '小科目': q['smallSubject'],
            'PDF页码': page,
            '题干摘要': qshort,
            '配图': ''
        })

print(f"\n需配图题目共 {len(img_questions)} 道")

# 验证PDF45页的题目
print("\nPDF45页的需配图题目:")
for q in img_questions:
    if q['PDF页码'] == '45':
        print(f"  id={q['id']}, 题号={q['题号']}, {q['题干摘要'][:40]}")

# 4. 生成Excel
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "需配图题目清单"

headers = ['序号', 'id', '题号', '年份', '小科目', 'PDF页码', '题干摘要', '配图']
ws.append(headers)

header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font_white = Font(bold=True, size=11, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font_white
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

for idx, q in enumerate(img_questions, start=1):
    row = [idx, q['id'], q['题号'], q['年份'], q['小科目'], q['PDF页码'], q['题干摘要'], q['配图']]
    ws.append(row)
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=idx+1, column=col)
        cell.border = thin_border
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if col in [1, 2, 4, 6]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 6
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 8
ws.column_dimensions['E'].width = 22
ws.column_dimensions['F'].width = 9
ws.column_dimensions['G'].width = 60
ws.column_dimensions['H'].width = 30
ws.freeze_panes = 'A2'
ws.row_dimensions[1].height = 25

output_path = r'D:\应用程序开发\刷题\需配图题目清单.xlsx'
wb.save(output_path)
print(f"\nExcel已保存: {output_path}")
